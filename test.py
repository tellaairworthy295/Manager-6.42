import json
from urllib.parse import urlencode

def main(inputs: str) -> dict:
    if isinstance(inputs, str):
        try:
            data = json.loads(inputs)
        except:
            data = {}
    else:
        data = inputs

    k_query = data.get("keywords", "")
    base = data.get("base", "")

    # Use urlencode to safely encode the query parameters
    params = {
        "q": f"site:{base}{k_query}",
        "tbs": "qdr:h,sbd:1",
        "tbm": "nws"
    }

    query_string = urlencode(params)
    bquery = f"https://www.google.com/search?{query_string}"

    return {"bloom_query": bquery}

def all_to_one(dir: str = "docx"):
    import os
    from docx import Document

    all_content = []

    # Iterate all .docx files in the directory
    for f in os.listdir(dir):
        if f.lower().endswith(".docx"):
            file_path = os.path.join(dir, f)
            try:
                doc = Document(file_path)
                text = "\n".join([para.text for para in doc.paragraphs])
                if text.strip():
                    all_content.append(text)
            except Exception as e:
                print(f"Failed to read {file_path}: {e}")

    # Combine all content into one docx file
    if all_content:
        combined_path = os.path.join(dir, "combined.docx")
        combined_doc = Document()
        for entry in all_content:
            combined_doc.add_paragraph(entry)
            combined_doc.add_paragraph()  # Blank line between articles
        combined_doc.save(combined_path)
    
if __name__ == "__main__":
    #print(main("\n\n\n\n  \n  \n  \n  \n\n\n{\n  \"base\": \"bloomberg.com/news/articles\",\n  \"keywords\": \" China OR \"Hong Kong\" OR AI\"\n}"))
    all_to_one()